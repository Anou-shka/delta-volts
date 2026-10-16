"""The Excel model builds, recalculates with zero errors, and every check is TRUE."""
import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / "model" / "delta_model.xlsx"

pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def recalculated():
    from src.delta import build_excel, recalc
    assert build_excel.main() == 0
    result = recalc.recalc(MODEL)
    assert result.get("status") == "success", result
    return load_workbook(MODEL, data_only=True), load_workbook(MODEL)


def test_all_checks_true(recalculated):
    values, _ = recalculated
    cover = values["Cover"]
    flag = next(r[2].value for r in cover.iter_rows() if r[1].value == "ALL CHECKS PASS")
    assert flag is True
    checks = values["Checks"]
    for row in checks.iter_rows(min_row=8):
        for c in row[2:]:
            assert c.value is not False, f"Checks!{c.coordinate} is FALSE"


@pytest.mark.parametrize("sheet,label,col,value", [
    ("IS", "Revenue", "G", 554_885.168),
    ("IS", "Operating profit (EBIT)", "G", 83_932.074),
    ("IS", "EPS, basic (NT$)", "J", 17.59),
    ("BS", "Total assets", "G", 639_618.600),
    ("CF", "Cash at end of period", "J", 182_350.204),
])
def test_key_values(recalculated, sheet, label, col, value):
    values, _ = recalculated
    ws = values[sheet]
    row = next(r[0].row for r in ws.iter_rows() if r[0].value == label)
    assert ws[f"{col}{row}"].value == pytest.approx(value, abs=0.01)


def test_statements_have_no_hardcoded_numbers(recalculated):
    """Inputs live on Historicals/Segments/Assumptions; statements and schedules are formulas only."""
    _, formulas = recalculated
    for sheet in ("IS", "BS", "CF", "Schedules"):
        ws = formulas[sheet]
        for row in ws.iter_rows(min_row=7, min_col=3, max_col=10):
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    pytest.fail(f"hard-coded number at {sheet}!{c.coordinate}: {c.value}")
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert not re.search(r"(?<![A-Z$\d.])\d{3,}", c.value), f"constant in {sheet}!{c.coordinate}"
