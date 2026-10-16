"""Recalculate an .xlsx in place with headless LibreOffice and report formula errors.

openpyxl writes formulas without cached values; this opens the workbook in
LibreOffice, runs calculateAll(), saves, then scans every cell for errors.

Run: python -m src.delta.recalc model/delta_model.xlsx
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")


def _env() -> dict:
    env = os.environ.copy()
    env.setdefault("SAL_USE_VCLPLUGIN", "svp")
    return env


def recalc(path: str | Path, timeout: int = 120) -> dict:
    path = Path(path).resolve()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"error": "LibreOffice (soffice) not found"}
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as prof:
        url = Path(prof).as_uri()
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=_env())
        mdir = Path(prof) / "user" / "basic" / "Standard"
        if not mdir.exists():
            return {"error": "LibreOffice profile not created"}
        (mdir / "Module1.xba").write_text(MACRO)
        before = path.stat().st_mtime_ns
        r = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
                            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                            str(path)], capture_output=True, text=True, timeout=timeout, env=_env())
        if r.returncode != 0 or path.stat().st_mtime_ns == before:
            return {"error": f"recalculation failed: {r.stderr.strip()[:500]}"}
    return scan(path)


def scan(path: str | Path) -> dict:
    """Count formulas and error values in a recalculated workbook."""
    wf = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    n_formulas, errors = 0, []
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formulas += 1
                    v = wv[ws.title][c.coordinate].value
                    if v is None or (isinstance(v, str) and v.startswith(ERRORS)):
                        errors.append(f"{ws.title}!{c.coordinate}: {v}")
    return {"status": "success" if not errors else "errors_found", "total_formulas": n_formulas,
            "total_errors": len(errors), "errors": errors[:50]}


if __name__ == "__main__":
    out = recalc(sys.argv[1])
    print(json.dumps(out, indent=2))
    sys.exit(0 if out.get("status") == "success" else 1)
