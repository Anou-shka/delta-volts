"""Build model/delta_model.xlsx with live formulas from data/processed/.

M2 scope: Cover, Assumptions (historical driver reference + scenario selector),
Historicals, Segments, IS, BS, CF, Schedules, Checks. Historical periods
FY2021A-FY2025A plus 1H25A and 1H26A. Forecast columns are added in M3.

Conventions (banking): blue = hard-coded input, black = formula, green = link to
another sheet, yellow fill = key input for the analyst. Units: NT$ millions
unless stated. Every hard-coded figure is in data/sources.csv.

Run: python -m src.delta.build_excel   (then python -m src.delta.recalc model/delta_model.xlsx)
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .extract import BS_MAP, CF_MAP, IS_MAP

ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "data" / "processed"
OUT = ROOT / "model" / "delta_model.xlsx"

FY = ["FY2021", "FY2022", "FY2023", "FY2024", "FY2025"]
INTERIM = ["1H25", "1H26"]
PERIODS = FY + INTERIM
COL = {"FY2021": "C", "FY2022": "D", "FY2023": "E", "FY2024": "F", "FY2025": "G", "1H25": "I", "1H26": "J"}
HEAD = {p: f"{p}A" for p in PERIODS}
PERIOD_END = {"FY2021": "31-Dec-21", "FY2022": "31-Dec-22", "FY2023": "31-Dec-23", "FY2024": "31-Dec-24",
              "FY2025": "31-Dec-25", "1H25": "30-Jun-25", "1H26": "30-Jun-26"}
PREV = {"FY2022": "FY2021", "FY2023": "FY2022", "FY2024": "FY2023", "FY2025": "FY2024", "1H26": "1H25"}
PREV_BS = {"FY2022": "FY2021", "FY2023": "FY2022", "FY2024": "FY2023", "FY2025": "FY2024",
           "1H25": "FY2024", "1H26": "FY2025"}  # balance sheet at the start of the period

# ---------------------------------------------------------------- styles
ARIAL = "Arial"
BLUE, BLACK, GREEN, GREY, WHITE = "0000FF", "000000", "008000", "595959", "FFFFFF"
NAVY = "1F3864"
F_NUM = '#,##0.0;(#,##0.0);"-"'
F_PCT = '0.0%;(0.0%);"-"'
F_EPS = '0.00;(0.00);"-"'
F_DAYS = '0.0;(0.0);"-"'
F_X = '0.00"x";(0.00"x");"-"'
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TRUE = PatternFill("solid", fgColor="C6EFCE")
FILL_FALSE = PatternFill("solid", fgColor="FFC7CE")
THIN_TOP = Border(top=Side(style="thin", color="808080"))
STYLE_COLOR = {"input": BLUE, "formula": BLACK, "link": GREEN}

LINK_RE = re.compile(r"^=-?'[^']+'![A-Z]+\d+$")
REG: dict[tuple[str, str], int] = {}  # (sheet, key) -> row


def font(color=BLACK, bold=False, italic=False, size=10):
    return Font(name=ARIAL, color=color, bold=bold, italic=italic, size=size)


def tf_format(ws, rng: str, top_left: str) -> None:
    """Green fill for TRUE, red for FALSE; blank cells stay unfilled."""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISLOGICAL({top_left}),{top_left})"], fill=FILL_TRUE))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISLOGICAL({top_left}),NOT({top_left}))"],
                                                   fill=FILL_FALSE))


def ref(sheet: str, key: str, p: str) -> str:
    return f"'{sheet}'!{COL[p]}{REG[(sheet, key)]}"


class Sheet:
    """Writes a period-grid sheet: labels in A, notes in B, periods in C..J."""

    def __init__(self, wb: Workbook, name: str, title: str, note: str, periods=PERIODS):
        self.ws = wb.create_sheet(name)
        self.name, self.periods, self.r = name, periods, 7
        ws = self.ws
        ws["A1"] = title
        ws["A1"].font = font(NAVY, bold=True, size=14)
        ws["A2"] = note
        ws["A2"].font = font(GREY, italic=True)
        for c in ("A3", "A4"):
            ws[c].font = font()
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 30
        for p in PERIODS:
            ws.column_dimensions[COL[p]].width = 12
        ws.column_dimensions["H"].width = 3
        ws["A5"], ws["B5"] = "NT$ million", "Note / source"
        for c in ("A5", "B5"):
            ws[c].font = font(WHITE, bold=True)
            ws[c].fill = FILL_HEAD
        for p in periods:
            h = ws[f"{COL[p]}5"]
            h.value = HEAD[p]
            h.font, h.fill, h.alignment = font(WHITE, bold=True), FILL_HEAD, Alignment(horizontal="right")
            d = ws[f"{COL[p]}6"]
            d.value = PERIOD_END[p]
            d.font, d.alignment = font(GREY, italic=True), Alignment(horizontal="right")
        if any(p in periods for p in INTERIM):
            ws["I4"] = "Interim (six months)"
            ws["I4"].font = font(GREY, italic=True)
        ws.freeze_panes = "C7"

    def section(self, text: str):
        self.r += 1
        for c in ["A", "B"] + [COL[p] for p in self.periods]:
            self.ws[f"{c}{self.r}"].fill = FILL_SECTION
        self.ws[f"A{self.r}"] = text
        self.ws[f"A{self.r}"].font = font(NAVY, bold=True)
        self.r += 1

    def line(self, key, label, fn, fmt=F_NUM, style="formula", bold=False, note="", periods=None, total=False):
        row = self.r
        REG[(self.name, key)] = row
        ws = self.ws
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = font(bold=bold or total)
        if note:
            ws[f"B{row}"] = note
            ws[f"B{row}"].font = font(GREY, italic=True, size=9)
        for p in periods or self.periods:
            v = fn(p)
            if v is None:
                continue
            c = ws[f"{COL[p]}{row}"]
            c.value = v
            color = STYLE_COLOR[style]
            if style == "link" and isinstance(v, str) and not LINK_RE.match(v):
                color = BLACK  # a formula combining links is a calculation, not a link
            c.font = font(color, bold=bold or total)
            c.number_format = fmt
            if total:
                c.border = THIN_TOP
        self.r += 1
        return row

    def ref(self, key, p):
        return f"{COL[p]}{REG[(self.name, key)]}"

    def blank(self, n=1):
        self.r += n


# ---------------------------------------------------------------- data
def load():
    h = pd.read_csv(PROCESSED / "historicals.csv", index_col=0)
    s = pd.read_csv(PROCESSED / "segments.csv")
    lines = pd.read_csv(PROCESSED / "fs_lines.csv")
    checks = pd.read_csv(PROCESSED / "checks.csv")
    return h, s, lines, checks


HIST_LABELS = {
    # IS
    "revenue": "Operating revenue", "cogs": "Operating costs", "gross_profit": "Gross profit (reported)",
    "selling": "Selling expenses", "g_and_a": "General and administrative expenses",
    "r_and_d": "Research and development expenses", "credit_loss": "Expected credit impairment (loss) gain",
    "opex_total": "Total operating expenses (reported)", "ebit": "Operating profit (reported)",
    "interest_income": "Interest income", "other_income": "Other income",
    "other_gains_losses": "Other gains and losses", "finance_costs": "Finance costs",
    "equity_method_income": "Share of profit (loss) of associates and JVs",
    "non_operating_total": "Total non-operating income and expenses (reported)",
    "pbt": "Profit before income tax (reported)", "tax": "Income tax expense",
    "net_income": "Profit for the period (reported)", "ni_parent": "Profit attributable to owners of the parent",
    "ni_nci": "Profit attributable to non-controlling interests", "eps_basic": "Basic EPS (NT$)",
    "eps_diluted": "Diluted EPS (NT$)",
    # BS
    "cash": "Cash and cash equivalents", "fvtpl_current": "Financial assets at FVTPL - current",
    "fvoci_current": "Financial assets at FVOCI - current",
    "amortised_cost_current": "Financial assets at amortised cost - current",
    "contract_assets_current": "Contract assets - current", "notes_receivable": "Notes receivable, net",
    "accounts_receivable": "Accounts receivable, net", "ar_related": "Accounts receivable - related parties",
    "inventories": "Inventories", "total_current_assets": "Total current assets (reported)",
    "fvtpl_noncurrent": "Financial assets at FVTPL - non-current",
    "fvoci_noncurrent": "Financial assets at FVOCI - non-current",
    "amortised_cost_noncurrent": "Financial assets at amortised cost - non-current",
    "equity_method_investments": "Investments accounted for under the equity method",
    "ppe": "Property, plant and equipment", "right_of_use": "Right-of-use assets",
    "investment_property": "Investment property, net", "intangibles": "Intangible assets (incl. goodwill)",
    "total_noncurrent_assets": "Total non-current assets (reported)", "total_assets": "Total assets (reported)",
    "st_borrowings": "Short-term borrowings", "notes_payable": "Notes payable", "accounts_payable": "Accounts payable",
    "ap_related": "Accounts payable - related parties", "lt_debt_current": "Long-term liabilities, current portion",
    "lease_liab_current": "Lease liabilities - current (if shown separately)",
    "total_current_liabilities": "Total current liabilities (reported)", "bonds_payable": "Bonds payable",
    "lt_borrowings": "Long-term borrowings", "lease_liab_noncurrent": "Lease liabilities - non-current",
    "total_noncurrent_liabilities": "Total non-current liabilities (reported)",
    "total_liabilities": "Total liabilities (reported)", "share_capital": "Share capital - common stock",
    "equity_parent": "Equity attributable to owners of the parent", "nci": "Non-controlling interest",
    "total_equity": "Total equity (reported)", "total_liabilities_equity": "Total liabilities and equity (reported)",
    # CF
    "cf_pbt": "Profit before tax (cash flow statement)", "depreciation": "Depreciation",
    "amortisation": "Amortisation", "cf_change_ar": "Change in accounts receivable",
    "cf_change_inventories": "Change in inventories", "cf_change_ap": "Change in accounts payable",
    "interest_paid": "Interest paid", "taxes_paid": "Income taxes paid",
    "cfo": "Net cash flows from operating activities (reported)",
    "capex_ppe": "Acquisition of property, plant and equipment",
    "proceeds_ppe": "Proceeds from disposal of PP&E", "capex_intangibles": "Acquisition of intangible assets",
    "acquisitions": "Acquisition of subsidiaries, net of cash acquired",
    "cfi": "Net cash flows from investing activities (reported)",
    "st_borrowings_change": "Increase (decrease) in short-term borrowings",
    "bonds_issued": "Issuance of bonds payable", "bonds_repaid": "Repayment of bonds payable",
    "lt_debt_proceeds": "Proceeds from long-term borrowings", "lt_debt_repaid": "Repayment of long-term borrowings",
    "lease_principal": "Lease principal repayment", "dividends_paid": "Cash dividends paid (owners of the parent)",
    "dividends_paid_nci": "Cash dividends paid to non-controlling interests",
    "cff": "Net cash flows from financing activities (reported)",
    "fx_effect": "Effect of exchange rate changes on cash",
    "net_change_cash": "Net increase (decrease) in cash (reported)",
    "cash_begin": "Cash at beginning of period (reported)", "cash_end": "Cash at end of period (reported)",
}
SKIP_BS_PERIODS = set()  # all historical columns carry a balance sheet (1H = 30 June)


# ---------------------------------------------------------------- sheets
def build_historicals(wb, h, checks):
    sh = Sheet(wb, "Historicals", "Historical financial statements - as reported",
               "NT$ million (reported in NT$ thousands, shown /1,000). Taiwan IFRS, consolidated. Blue = hard-coded "
               "from the filings; page reference for every value in data/sources.csv.")
    ws = sh.ws
    docs = {**{p: f"fs_{p[2:]}_annual_consolidated.pdf" for p in FY},
            "1H25": "fs_2026q2_consolidated.pdf (comparative)", "1H26": "fs_2026q2_consolidated.pdf"}
    ws["A3"] = "Source document"
    ws["A3"].font = font(GREY, italic=True)
    for p in PERIODS:
        ws[f"{COL[p]}3"] = docs[p].replace("_consolidated.pdf", "").replace("fs_", "")
        ws[f"{COL[p]}3"].font = font(GREY, italic=True, size=8)
    for title, mp in (("Income statement", IS_MAP), ("Balance sheet (period end)", BS_MAP),
                      ("Cash flow statement", CF_MAP)):
        sh.section(title)
        for key in mp:
            if key not in h.index:
                continue
            vals = h.loc[key]
            fmt = F_EPS if key.startswith("eps") else F_NUM
            sh.line(key, HIST_LABELS.get(key, key), lambda p, v=vals: (None if pd.isna(v[p]) else float(v[p])),
                    fmt=fmt, style="input", note="extracted; see sources.csv")
    sh.section("Other inputs")
    sh.line("par_value", "Par value per share (NT$)", lambda p: 10.0, fmt=F_EPS, style="input",
            note="Taiwan standard par; 2,598m shares per ic_4Q25 p.10")
    sh.line("checks_total", "Python extraction tie-outs: total", lambda p: float(len(checks)), fmt="0",
            style="input", note="data/processed/checks.csv", periods=["FY2021"])
    sh.line("checks_passed", "Python extraction tie-outs: passed", lambda p: float(checks["pass"].sum()), fmt="0",
            style="input", note="data/processed/checks.csv", periods=["FY2021"])
    return sh


def H(key, p):
    return ref("Historicals", key, p)


def build_is(wb):
    sh = Sheet(wb, "IS", "Income statement", "NT$ million unless stated. Lines link to Historicals; subtotals are "
               "recomputed (not copied) and tested against reported figures on the Checks sheet.")
    lk = lambda k: (lambda p: f"={H(k, p)}")  # noqa: E731
    sh.section("Income statement")
    sh.line("revenue", "Revenue", lk("revenue"), style="link", bold=True)
    sh.line("cogs", "Cost of revenue", lk("cogs"), style="link")
    sh.line("gross_profit", "Gross profit", lambda p: f"={sh.ref('revenue', p)}+{sh.ref('cogs', p)}", total=True)
    sh.line("selling", "Selling expenses", lk("selling"), style="link")
    sh.line("g_and_a", "General and administrative", lk("g_and_a"), style="link")
    sh.line("r_and_d", "Research and development", lk("r_and_d"), style="link")
    sh.line("credit_loss", "Expected credit loss", lk("credit_loss"), style="link")
    sh.line("opex", "Total operating expenses",
            lambda p: f"=SUM({sh.ref('selling', p)}:{sh.ref('credit_loss', p)})", total=True)
    sh.line("ebit", "Operating profit (EBIT)", lambda p: f"={sh.ref('gross_profit', p)}+{sh.ref('opex', p)}",
            total=True)
    sh.line("interest_income", "Interest income", lk("interest_income"), style="link")
    sh.line("finance_costs", "Finance costs", lk("finance_costs"), style="link")
    sh.line("other_nonop", "Other non-operating items",
            lambda p: f"={H('other_income', p)}+{H('other_gains_losses', p)}+{H('equity_method_income', p)}",
            note="other income, gains/losses, associates")
    sh.line("nonop", "Total non-operating income",
            lambda p: f"=SUM({sh.ref('interest_income', p)}:{sh.ref('other_nonop', p)})", total=True)
    sh.line("pbt", "Profit before tax", lambda p: f"={sh.ref('ebit', p)}+{sh.ref('nonop', p)}", total=True)
    sh.line("tax", "Income tax", lk("tax"), style="link")
    sh.line("net_income", "Net income", lambda p: f"={sh.ref('pbt', p)}+{sh.ref('tax', p)}", total=True)
    sh.line("ni_nci", "Less: non-controlling interests", lambda p: f"=-{H('ni_nci', p)}", style="link")
    sh.line("ni_parent", "Net income to owners of the parent",
            lambda p: f"={sh.ref('net_income', p)}+{sh.ref('ni_nci', p)}", total=True)
    sh.blank()
    sh.line("shares", "Shares outstanding, period end (m)",
            lambda p: f"={H('share_capital', p)}/{H('par_value', p)}",
            note="share capital / NT$10 par")
    sh.line("eps", "EPS, basic (NT$)", lambda p: f"={sh.ref('ni_parent', p)}/{sh.ref('shares', p)}", fmt=F_EPS,
            bold=True, note="computed; reported EPS on Historicals")
    sh.section("Memo")
    sh.line("d_and_a", "Depreciation and amortisation",
            lambda p: f"={H('depreciation', p)}+{H('amortisation', p)}", note="from cash flow statement")
    sh.line("ebitda", "EBITDA (EBIT + D&A)", lambda p: f"={sh.ref('ebit', p)}+{sh.ref('d_and_a', p)}", bold=True,
            note="operating definition (decision D3)")
    sh.section("Ratios")
    growth = lambda k: (lambda p: f"={sh.ref(k, p)}/{sh.ref(k, PREV[p])}-1" if p in PREV else None)  # noqa: E731
    sh.line("rev_growth", "Revenue growth (y/y)", growth("revenue"), fmt=F_PCT,
            note="1H26 vs 1H25")
    pct = lambda k, sign="": (lambda p: f"={sign}{sh.ref(k, p)}/{sh.ref('revenue', p)}")  # noqa: E731
    sh.line("gm", "Gross margin", pct("gross_profit"), fmt=F_PCT)
    sh.line("rnd_pct", "R&D % of revenue", pct("r_and_d", "-"), fmt=F_PCT)
    sh.line("sga_pct", "SG&A % of revenue",
            lambda p: f"=-({sh.ref('selling', p)}+{sh.ref('g_and_a', p)})/{sh.ref('revenue', p)}", fmt=F_PCT)
    sh.line("ebit_margin", "EBIT margin", pct("ebit"), fmt=F_PCT, bold=True)
    sh.line("ebitda_margin", "EBITDA margin", pct("ebitda"), fmt=F_PCT)
    sh.line("etr", "Effective tax rate", lambda p: f"=-{sh.ref('tax', p)}/{sh.ref('pbt', p)}", fmt=F_PCT)
    sh.line("ni_margin", "Net margin (to parent)", pct("ni_parent"), fmt=F_PCT)
    sh.line("eps_growth", "EPS growth (y/y)", growth("eps"), fmt=F_PCT)
    return sh


def build_bs(wb):
    sh = Sheet(wb, "BS", "Balance sheet", "NT$ million, period end. 'Other' lines are the reported subtotal less "
               "the lines shown, so each section ties to the reported figure; the Checks sheet tests that total "
               "assets equal total liabilities plus equity.")
    lk = lambda k: (lambda p: f"={H(k, p)}")  # noqa: E731
    s = lambda *ks: (lambda p: "=" + "+".join(H(k, p) for k in ks))  # noqa: E731
    sh.section("Assets")
    sh.line("cash", "Cash and cash equivalents", lk("cash"), style="link")
    sh.line("st_investments", "Short-term investments", s("fvtpl_current", "fvoci_current", "amortised_cost_current"),
            note="FVTPL + FVOCI + amortised cost")
    sh.line("receivables", "Receivables", s("notes_receivable", "accounts_receivable", "ar_related"),
            note="notes + trade + related parties")
    sh.line("inventories", "Inventories", lk("inventories"), style="link")
    sh.line("other_ca", "Other current assets",
            lambda p: f"={H('total_current_assets', p)}-SUM({sh.ref('cash', p)}:{sh.ref('inventories', p)})",
            note="incl. contract assets, prepayments, tax")
    sh.line("tca", "Total current assets", lambda p: f"=SUM({sh.ref('cash', p)}:{sh.ref('other_ca', p)})",
            total=True)
    sh.line("lt_investments", "Long-term investments",
            s("fvtpl_noncurrent", "fvoci_noncurrent", "amortised_cost_noncurrent", "equity_method_investments"))
    sh.line("ppe", "Property, plant and equipment", lk("ppe"), style="link")
    sh.line("rou", "Right-of-use assets", lk("right_of_use"), style="link")
    sh.line("intangibles", "Intangible assets (incl. goodwill)", lk("intangibles"), style="link")
    sh.line("other_nca", "Other non-current assets",
            lambda p: f"={H('total_noncurrent_assets', p)}-SUM({sh.ref('lt_investments', p)}:{sh.ref('intangibles', p)})",
            note="incl. deferred tax, investment property")
    sh.line("tnca", "Total non-current assets",
            lambda p: f"=SUM({sh.ref('lt_investments', p)}:{sh.ref('other_nca', p)})", total=True)
    sh.line("ta", "Total assets", lambda p: f"={sh.ref('tca', p)}+{sh.ref('tnca', p)}", bold=True, total=True)
    sh.section("Liabilities and equity")
    sh.line("st_debt", "Short-term debt", s("st_borrowings", "lt_debt_current"),
            note="ST borrowings + current portion of LT debt")
    sh.line("payables", "Payables", s("notes_payable", "accounts_payable", "ap_related"),
            note="notes + trade + related parties")
    sh.line("other_cl", "Other current liabilities",
            lambda p: f"={H('total_current_liabilities', p)}-{sh.ref('st_debt', p)}-{sh.ref('payables', p)}",
            note="incl. other payables, tax, contract liab., leases")
    sh.line("tcl", "Total current liabilities", lambda p: f"=SUM({sh.ref('st_debt', p)}:{sh.ref('other_cl', p)})",
            total=True)
    sh.line("lt_debt", "Long-term debt", s("bonds_payable", "lt_borrowings"), note="bonds + LT borrowings")
    sh.line("leases", "Lease liabilities (non-current)", lk("lease_liab_noncurrent"), style="link")
    sh.line("other_ncl", "Other non-current liabilities",
            lambda p: f"={H('total_noncurrent_liabilities', p)}-{sh.ref('lt_debt', p)}-{sh.ref('leases', p)}",
            note="incl. deferred tax liabilities")
    sh.line("tncl", "Total non-current liabilities",
            lambda p: f"=SUM({sh.ref('lt_debt', p)}:{sh.ref('other_ncl', p)})", total=True)
    sh.line("tl", "Total liabilities", lambda p: f"={sh.ref('tcl', p)}+{sh.ref('tncl', p)}", total=True)
    sh.line("equity_parent", "Equity attributable to owners of the parent", lk("equity_parent"), style="link")
    sh.line("nci", "Non-controlling interest", lk("nci"), style="link")
    sh.line("te", "Total equity", lambda p: f"={sh.ref('equity_parent', p)}+{sh.ref('nci', p)}", total=True)
    sh.line("tle", "Total liabilities and equity", lambda p: f"={sh.ref('tl', p)}+{sh.ref('te', p)}", bold=True,
            total=True)
    sh.line("bs_diff", "Balance check: assets - (liabilities + equity)",
            lambda p: f"={sh.ref('ta', p)}-{sh.ref('tle', p)}", note="must be zero")
    sh.section("Memo")
    sh.line("total_debt", "Total debt", lambda p: f"={sh.ref('st_debt', p)}+{sh.ref('lt_debt', p)}", bold=True)
    sh.line("net_debt", "Net debt / (net cash)",
            lambda p: f"={sh.ref('total_debt', p)}-{sh.ref('cash', p)}-{sh.ref('st_investments', p)}",
            note="debt - cash - ST investments; excl. leases")
    sh.line("nwc", "Trade working capital",
            lambda p: f"={sh.ref('receivables', p)}+{sh.ref('inventories', p)}-{sh.ref('payables', p)}",
            note="receivables + inventories - payables")
    return sh


def build_cf(wb, bs):
    sh = Sheet(wb, "CF", "Cash flow statement", "NT$ million. Indirect method as reported (Taiwan IFRS: interest "
               "and tax paid sit in operating cash flow). 'Other' lines tie each section to the reported total; "
               "closing cash is tested against the balance sheet on the Checks sheet.")
    lk = lambda k: (lambda p: f"={H(k, p)}")  # noqa: E731
    sh.section("Operating activities")
    sh.line("pbt", "Profit before tax", lk("cf_pbt"), style="link")
    sh.line("d_and_a", "Depreciation and amortisation", lambda p: f"={H('depreciation', p)}+{H('amortisation', p)}")
    sh.line("chg_ar", "(Increase) decrease in accounts receivable", lk("cf_change_ar"), style="link")
    sh.line("chg_inv", "(Increase) decrease in inventories", lk("cf_change_inventories"), style="link")
    sh.line("chg_ap", "Increase (decrease) in accounts payable", lk("cf_change_ap"), style="link")
    sh.line("interest_paid", "Interest paid", lk("interest_paid"), style="link")
    sh.line("taxes_paid", "Income taxes paid", lk("taxes_paid"), style="link")
    sh.line("other_op", "Other operating items",
            lambda p: f"={H('cfo', p)}-SUM({sh.ref('pbt', p)}:{sh.ref('taxes_paid', p)})",
            note="non-cash items, other WC, interest received")
    sh.line("cfo", "Cash flow from operations", lambda p: f"=SUM({sh.ref('pbt', p)}:{sh.ref('other_op', p)})",
            total=True)
    sh.section("Investing activities")
    sh.line("capex_ppe", "Purchase of PP&E", lk("capex_ppe"), style="link")
    sh.line("proceeds_ppe", "Disposal of PP&E", lk("proceeds_ppe"), style="link")
    sh.line("capex_int", "Purchase of intangible assets", lk("capex_intangibles"), style="link")
    sh.line("acquisitions", "Acquisitions, net of cash acquired", lk("acquisitions"), style="link")
    sh.line("other_inv", "Other investing items",
            lambda p: f"={H('cfi', p)}-SUM({sh.ref('capex_ppe', p)}:{sh.ref('acquisitions', p)})",
            note="financial assets, prepayments, other")
    sh.line("cfi", "Cash flow from investing", lambda p: f"=SUM({sh.ref('capex_ppe', p)}:{sh.ref('other_inv', p)})",
            total=True)
    sh.section("Financing activities")
    sh.line("net_borrowing", "Net borrowing / (repayment)",
            lambda p: "=" + "+".join(H(k, p) for k in ("st_borrowings_change", "bonds_issued", "bonds_repaid",
                                                       "lt_debt_proceeds", "lt_debt_repaid")),
            note="ST + bonds + LT borrowings, net")
    sh.line("lease_principal", "Lease principal repaid", lk("lease_principal"), style="link")
    sh.line("dividends", "Dividends paid to shareholders", lk("dividends_paid"), style="link",
            note="paid in Q3; nil in 1H")
    sh.line("dividends_nci", "Dividends paid to non-controlling interests", lk("dividends_paid_nci"), style="link")
    sh.line("other_fin", "Other financing items",
            lambda p: f"={H('cff', p)}-SUM({sh.ref('net_borrowing', p)}:{sh.ref('dividends_nci', p)})",
            note="deposits, NCI transactions")
    sh.line("cff", "Cash flow from financing",
            lambda p: f"=SUM({sh.ref('net_borrowing', p)}:{sh.ref('other_fin', p)})", total=True)
    sh.blank()
    sh.line("fx", "Effect of exchange rates", lk("fx_effect"), style="link")
    sh.line("net_change", "Net change in cash",
            lambda p: f"={sh.ref('cfo', p)}+{sh.ref('cfi', p)}+{sh.ref('cff', p)}+{sh.ref('fx', p)}", total=True)
    sh.line("cash_open", "Cash at beginning of period",
            lambda p: f"='BS'!{COL[PREV_BS[p]]}{REG[('BS', 'cash')]}" if p in PREV_BS else f"={H('cash_begin', p)}",
            style="link", note="prior balance sheet; FY2021 from CF")
    sh.line("cash_close", "Cash at end of period",
            lambda p: f"={sh.ref('cash_open', p)}+{sh.ref('net_change', p)}", bold=True, total=True)
    sh.section("Memo")
    sh.line("capex", "Capex (PP&E + intangibles)",
            lambda p: f"={sh.ref('capex_ppe', p)}+{sh.ref('capex_int', p)}", note="decision D4")
    sh.line("fcf", "Free cash flow (CFO + capex)", lambda p: f"={sh.ref('cfo', p)}+{sh.ref('capex', p)}", bold=True)
    sh.line("fcf_margin", "FCF margin", lambda p: f"={sh.ref('fcf', p)}/'IS'!{COL[p]}{REG[('IS', 'revenue')]}",
            fmt=F_PCT)
    sh.line("cash_conversion", "Cash conversion (CFO / net income)",
            lambda p: f"={sh.ref('cfo', p)}/'IS'!{COL[p]}{REG[('IS', 'net_income')]}", fmt=F_X)
    return sh


SEGS = ["Power Electronics", "Infrastructure", "Automation", "Mobility"]
SEG_KEY = {"Power Electronics": "pe", "Infrastructure": "infra", "Automation": "auto", "Mobility": "mob"}


def build_segments(wb, segs):
    sh = Sheet(wb, "Segments", "Segment revenue and profit",
               "NT$ million. Source: operating segment note, latest restated basis. FY2021-22 are on Delta's old "
               "3-segment basis (Mobility inside Power Electronics), so FY2023 growth is not shown for Power Electronics or "
               "Mobility. Infrastructure and Automation were restated by <0.3% and are shown.")
    latest = segs[segs.latest]
    get = lambda seg, item, p: latest[(latest.segment == seg) & (latest["item"] == item) & (latest.period == p)]  # noqa
    IS = lambda k, p: f"'IS'!{COL[p]}{REG[('IS', k)]}"  # noqa: E731

    def val(seg, item):
        def f(p):
            r = get(seg, item, p)
            return None if r.empty else float(r.value.iloc[0]) / 1000
        return f

    sh.section("Revenue from external customers")
    for seg in SEGS:
        sh.line(f"rev_{SEG_KEY[seg]}", seg, val(seg, "revenue"), style="input",
                note="in PE before FY2023" if seg == "Mobility" else "segment note")
    sh.line("rev_others", "Others / unallocated",
            lambda p: f"={IS('revenue', p)}-SUM({sh.ref('rev_pe', p)}:{sh.ref('rev_mob', p)})",
            note="group revenue - segments ('Others')")
    sh.line("rev_total", "Group revenue", lambda p: f"=SUM({sh.ref('rev_pe', p)}:{sh.ref('rev_others', p)})",
            total=True)
    sh.section("Segment operating profit")
    for seg in SEGS:
        sh.line(f"op_{SEG_KEY[seg]}", seg, val(seg, "op_profit"), style="input", note="segment note")
    sh.line("op_total", "Total segment profit", lambda p: f"=SUM({sh.ref('op_pe', p)}:{sh.ref('op_mob', p)})",
            total=True, note="equals EBIT (tested on Checks)")
    sh.section("Revenue growth (y/y)")
    comparable = {"FY2022": SEGS[:3], "FY2023": ["Infrastructure", "Automation"], "FY2024": SEGS,
                  "FY2025": SEGS, "1H26": SEGS}
    for seg in SEGS:
        k = SEG_KEY[seg]
        sh.line(f"g_{k}", seg, lambda p, s=seg, k=k: (f"={sh.ref('rev_' + k, p)}/{sh.ref('rev_' + k, PREV[p])}-1"
                                                      if s in comparable.get(p, []) else None), fmt=F_PCT)
    sh.line("g_total", "Group", lambda p: f"={sh.ref('rev_total', p)}/{sh.ref('rev_total', PREV[p])}-1"
            if p in PREV else None, fmt=F_PCT, bold=True)
    sh.section("Revenue mix (% of group)")
    for seg in SEGS:
        k = SEG_KEY[seg]
        sh.line(f"mix_{k}", seg, lambda p, k=k, s=seg: (None if s == "Mobility" and p in ("FY2021", "FY2022") else
                                                        f"={sh.ref('rev_' + k, p)}/{sh.ref('rev_total', p)}"), fmt=F_PCT)
    sh.section("Segment margin (profit / revenue)")
    for seg in SEGS:
        k = SEG_KEY[seg]
        sh.line(f"m_{k}", seg, lambda p, k=k, s=seg: (None if s == "Mobility" and p in ("FY2021", "FY2022") else
                                                      f"={sh.ref('op_' + k, p)}/{sh.ref('rev_' + k, p)}"), fmt=F_PCT)
    sh.line("m_total", "Group EBIT margin", lambda p: f"={sh.ref('op_total', p)}/{sh.ref('rev_total', p)}",
            fmt=F_PCT, bold=True)

    # AI anchors (decision D1): Delta's own disclosures, used to build the AI line in M3
    ws = sh.ws
    sh.blank()
    sh.section("AI data-centre anchors (management disclosures, decision D1) - the AI line is built from these in M3")
    hdr = ["Anchor", "Source", "Period", "Share of revenue", "Qualifier", "Implied NT$ m"]
    cols = ["A", "B", "C", "D", "E", "F"]
    for c, t in zip(cols, hdr):
        ws[f"{c}{sh.r}"] = t
        ws[f"{c}{sh.r}"].font = font(WHITE, bold=True)
        ws[f"{c}{sh.r}"].fill = FILL_HEAD
    sh.r += 1
    anchors = [
        ("Liquid-cooling products", "2Q26 transcript p.12", "FY2025", 0.10, "approx.", "FY2025"),
        ("Liquid-cooling products", "2Q26 transcript p.12", "FY2026E", 0.12, "more than", None),
        ("AI-related products", "2Q26 transcript p.12", "FY2026E", 0.25, "more than", None),
        ("Data-centre-related business", "2Q26 transcript p.2", "1H26", 0.50, "more than", "1H26"),
    ]
    for name, src, per, share, qual, base in anchors:
        r = sh.r
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = name, src, per, share, qual
        ws[f"D{r}"].font, ws[f"D{r}"].number_format = font(BLUE), F_PCT
        ws[f"D{r}"].fill = FILL_KEY
        ws[f"D{r}"].comment = Comment(f"Management statement, {src} (Delta 2Q26 results meeting, 30 Jul 2026). "
                                      "Rounded; translated transcript (Mandarin prevails).", "model")
        for c in ("A", "B", "C", "E"):
            ws[f"{c}{r}"].font = font(size=9)
        if base:
            ws[f"F{r}"] = f"=D{r}*{IS('revenue', base)}"
            ws[f"F{r}"].number_format = F_NUM
            ws[f"F{r}"].font = font()
        else:
            ws[f"F{r}"] = "M3 forecast"
            ws[f"F{r}"].font = font(GREY, italic=True, size=9)
        REG[("Segments", f"anchor_{per}_{name[:6]}")] = r
        sh.r += 1
    ws[f"A{sh.r}"] = ("Not used: press-reported split (AI power >20%, AI cooling ~9% of FY2025 revenue) - not in "
                      "Delta's own documents; TODO-SOURCE in sources.csv.")
    ws[f"A{sh.r}"].font = font(GREY, italic=True, size=9)
    return sh


def build_schedules(wb):
    sh = Sheet(wb, "Schedules", "Supporting schedules (historical)",
               "NT$ million unless stated. Fiscal years only; ratios use year-end balances. These historical ratios "
               "anchor the forecast drivers set on the Assumptions sheet in M3.", periods=FY)
    ws = sh.ws
    IS = lambda k, p: f"'IS'!{COL[p]}{REG[('IS', k)]}"  # noqa: E731
    BS = lambda k, p: f"'BS'!{COL[p]}{REG[('BS', k)]}"  # noqa: E731
    CF = lambda k, p: f"'CF'!{COL[p]}{REG[('CF', k)]}"  # noqa: E731
    ws["A4"], ws["B4"] = "Days in year", 365
    ws["B4"].font, ws["B4"].fill = font(BLUE), FILL_KEY
    days = "$B$4"
    has_prev = lambda f: (lambda p: f(p) if p in PREV else None)  # noqa: E731

    sh.section("Working capital")
    sh.line("revenue", "Revenue", lambda p: f"={IS('revenue', p)}", style="link")
    sh.line("cogs", "Cost of revenue (positive)", lambda p: f"=-{IS('cogs', p)}", style="link")
    sh.line("receivables", "Receivables", lambda p: f"={BS('receivables', p)}", style="link")
    sh.line("inventories", "Inventories", lambda p: f"={BS('inventories', p)}", style="link")
    sh.line("payables", "Payables", lambda p: f"={BS('payables', p)}", style="link")
    sh.line("dso", "Receivable days (DSO)",
            lambda p: f"={sh.ref('receivables', p)}/{sh.ref('revenue', p)}*{days}", fmt=F_DAYS)
    sh.line("dio", "Inventory days (DIO)",
            lambda p: f"={sh.ref('inventories', p)}/{sh.ref('cogs', p)}*{days}", fmt=F_DAYS)
    sh.line("dpo", "Payable days (DPO)", lambda p: f"={sh.ref('payables', p)}/{sh.ref('cogs', p)}*{days}",
            fmt=F_DAYS)
    sh.line("ccc", "Cash conversion cycle (days)",
            lambda p: f"={sh.ref('dso', p)}+{sh.ref('dio', p)}-{sh.ref('dpo', p)}", fmt=F_DAYS, bold=True)
    sh.line("other_ca_pct", "Other current assets % revenue",
            lambda p: f"={BS('other_ca', p)}/{sh.ref('revenue', p)}", fmt=F_PCT)
    sh.line("other_cl_pct", "Other current liabilities % revenue",
            lambda p: f"={BS('other_cl', p)}/{sh.ref('revenue', p)}", fmt=F_PCT)

    sh.section("PP&E roll-forward")
    sh.line("ppe_open", "Opening PP&E", has_prev(lambda p: f"={BS('ppe', PREV[p])}"), style="link")
    sh.line("capex_ppe", "Add: purchases of PP&E", lambda p: f"=-{CF('capex_ppe', p)}", style="link")
    sh.line("dep", "Less: depreciation", lambda p: f"=-'Historicals'!{COL[p]}{REG[('Historicals', 'depreciation')]}",
            style="link", note="all depreciation incl. right-of-use")
    sh.line("ppe_other", "Other movements (plug)",
            has_prev(lambda p: f"={BS('ppe', p)}-{sh.ref('ppe_open', p)}-{sh.ref('capex_ppe', p)}"
                               f"-{sh.ref('dep', p)}"),
            note="disposals, FX, M&A, add-back of ROU depreciation")
    sh.line("ppe_close", "Closing PP&E", lambda p: f"={BS('ppe', p)}", style="link", total=True)
    sh.line("capex_pct", "PP&E capex % revenue", lambda p: f"={sh.ref('capex_ppe', p)}/{sh.ref('revenue', p)}",
            fmt=F_PCT)
    sh.line("capex_total_pct", "Total capex (PP&E + intangibles) % revenue",
            lambda p: f"=-{CF('capex', p)}/{sh.ref('revenue', p)}", fmt=F_PCT, bold=True)
    sh.line("dep_rate", "Depreciation % opening PP&E",
            has_prev(lambda p: f"=-{sh.ref('dep', p)}/{sh.ref('ppe_open', p)}"), fmt=F_PCT)
    sh.line("da_pct", "D&A % revenue", lambda p: f"={IS('d_and_a', p)}/{sh.ref('revenue', p)}", fmt=F_PCT)
    sh.line("capex_da", "Capex / D&A", lambda p: f"=-{CF('capex', p)}/{IS('d_and_a', p)}", fmt=F_X)

    sh.section("Intangibles")
    sh.line("amort", "Amortisation", lambda p: f"='Historicals'!{COL[p]}{REG[('Historicals', 'amortisation')]}",
            style="link")
    sh.line("capex_int", "Purchases of intangibles", lambda p: f"=-{CF('capex_int', p)}", style="link")
    sh.line("intangibles", "Intangible assets (incl. goodwill)", lambda p: f"={BS('intangibles', p)}", style="link")

    sh.section("Debt and interest")
    sh.line("debt", "Total debt", lambda p: f"={BS('total_debt', p)}", style="link")
    sh.line("avg_debt", "Average debt", has_prev(lambda p: f"=AVERAGE({sh.ref('debt', PREV[p])},{sh.ref('debt', p)})"))
    sh.line("fin_cost", "Finance costs", lambda p: f"=-{IS('finance_costs', p)}", style="link",
            note="incl. lease interest")
    sh.line("cost_debt", "Implied cost of debt", has_prev(lambda p: f"={sh.ref('fin_cost', p)}/{sh.ref('avg_debt', p)}"),
            fmt=F_PCT, bold=True)
    sh.line("cash", "Cash and cash equivalents", lambda p: f"={BS('cash', p)}", style="link")
    sh.line("avg_cash", "Average cash", has_prev(lambda p: f"=AVERAGE({sh.ref('cash', PREV[p])},{sh.ref('cash', p)})"))
    sh.line("int_inc", "Interest income", lambda p: f"={IS('interest_income', p)}", style="link")
    sh.line("cash_yield", "Yield on average cash", has_prev(lambda p: f"={sh.ref('int_inc', p)}/{sh.ref('avg_cash', p)}"),
            fmt=F_PCT)

    sh.section("Tax")
    sh.line("etr", "Effective tax rate", lambda p: f"={IS('etr', p)}", style="link", fmt=F_PCT)
    sh.line("cash_tax", "Cash taxes paid % PBT", lambda p: f"=-{CF('taxes_paid', p)}/{IS('pbt', p)}", fmt=F_PCT)

    sh.section("Shares, dividends and minorities")
    sh.line("shares", "Shares outstanding (m)", lambda p: f"={IS('shares', p)}", style="link", fmt=F_NUM)
    sh.line("div_paid", "Dividends paid in year", lambda p: f"=-{CF('dividends', p)}", style="link")
    sh.line("dps_paid", "Dividend per share paid in year (NT$)",
            lambda p: f"={sh.ref('div_paid', p)}/{sh.ref('shares', p)}", fmt=F_EPS)
    sh.line("payout", "Payout of prior-year net income",
            has_prev(lambda p: f"={sh.ref('div_paid', p)}/{IS('ni_parent', PREV[p])}"), fmt=F_PCT, bold=True,
            note="Taiwan: dividend paid in year after earned")
    sh.line("nci_share", "Non-controlling interests % net income",
            lambda p: f"=-{IS('ni_nci', p)}/{IS('net_income', p)}", fmt=F_PCT)

    sh.section("Returns")
    sh.line("roe", "Return on average equity (parent)",
            has_prev(lambda p: f"={IS('ni_parent', p)}/AVERAGE({BS('equity_parent', PREV[p])},{BS('equity_parent', p)})"),
            fmt=F_PCT, bold=True)
    sh.line("asset_turn", "Revenue / average total assets",
            has_prev(lambda p: f"={sh.ref('revenue', p)}/AVERAGE({BS('ta', PREV[p])},{BS('ta', p)})"), fmt=F_X)
    return sh


def build_assumptions(wb):
    sh = Sheet(wb, "Assumptions", "Assumptions",
               "All forecast drivers live on this sheet. M2: scenario selector and the historical reference for each "
               "driver. M3 adds FY2026E-FY2030E inputs with Bear / Base / Bull columns, set by the analyst.",
               periods=FY)
    ws = sh.ws
    ws["A3"], ws["B3"] = "Scenario (select)", "Base"
    ws["B3"].font, ws["B3"].fill = font(BLUE, bold=True), FILL_KEY
    dv = DataValidation(type="list", formula1='"Bear,Base,Bull"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws["A4"], ws["B4"] = "Scenario index (1 Bear, 2 Base, 3 Bull)", '=MATCH(B3,{"Bear","Base","Bull"},0)'
    ws["B4"].font = font()
    REG[("Assumptions", "scenario")] = 3
    S = lambda sheet, k, p: f"='{sheet}'!{COL[p]}{REG[(sheet, k)]}"  # noqa: E731
    ws.column_dimensions["L"].width = 12
    ws["L5"] = "Average"
    ws["L5"].font, ws["L5"].fill = font(WHITE, bold=True), FILL_HEAD
    ws["L6"] = "FY23-25 avail."
    ws["L6"].font = font(GREY, italic=True)
    rows = [
        ("Revenue growth", [
            ("Power Electronics", "Segments", "g_pe", F_PCT), ("Infrastructure", "Segments", "g_infra", F_PCT),
            ("Automation", "Segments", "g_auto", F_PCT), ("Mobility", "Segments", "g_mob", F_PCT),
            ("Group", "IS", "rev_growth", F_PCT)]),
        ("Margins and costs", [
            ("Gross margin", "IS", "gm", F_PCT), ("R&D % revenue", "IS", "rnd_pct", F_PCT),
            ("SG&A % revenue", "IS", "sga_pct", F_PCT), ("EBIT margin", "IS", "ebit_margin", F_PCT),
            ("Effective tax rate", "IS", "etr", F_PCT), ("NCI % net income", "Schedules", "nci_share", F_PCT)]),
        ("Working capital (days)", [
            ("DSO", "Schedules", "dso", F_DAYS), ("DIO", "Schedules", "dio", F_DAYS),
            ("DPO", "Schedules", "dpo", F_DAYS)]),
        ("Capex, D&A and financing", [
            ("Total capex % revenue", "Schedules", "capex_total_pct", F_PCT),
            ("Depreciation % opening PP&E", "Schedules", "dep_rate", F_PCT),
            ("D&A % revenue", "Schedules", "da_pct", F_PCT), ("Cost of debt", "Schedules", "cost_debt", F_PCT),
            ("Yield on cash", "Schedules", "cash_yield", F_PCT), ("Dividend payout", "Schedules", "payout", F_PCT)]),
    ]
    sh.r = 8
    for title, items in rows:
        sh.section(title + " - historical reference")
        for label, sheet, key, fmt in items:
            row = sh.line(f"ref_{key}", label, lambda p, s=sheet, k=key: (
                S(s, k, p) if wb[s][f"{COL[p]}{REG[(s, k)]}"].value is not None else None),
                fmt=fmt, style="link", note=f"{sheet} sheet")
            ws[f"L{row}"] = f"=AVERAGE(E{row}:G{row})"
            ws[f"L{row}"].number_format, ws[f"L{row}"].font = fmt, font()
    sh.blank()
    ws[f"A{sh.r}"] = ("Forecast inputs (FY2026E-FY2030E, Bear / Base / Bull) are added in M3. Base period moves to 9M26 "
                      "when Q3 2026 results are published (decision D2).")
    ws[f"A{sh.r}"].font = font(GREY, italic=True)
    return sh


def build_checks(wb):
    sh = Sheet(wb, "Checks", "Model integrity checks",
               "Every check must be TRUE. Tolerance in NT$ million (B3) and NT$ per share (B4). The overall result "
               "rolls up to the Cover sheet.")
    ws = sh.ws
    ws["A3"], ws["B3"] = "Tolerance (NT$ m)", 0.01
    ws["A4"], ws["B4"] = "EPS tolerance (NT$)", 0.02
    for c in ("B3", "B4"):
        ws[c].font = font(BLUE)
    ws.column_dimensions["L"].width = 12
    ws["L5"] = "All periods"
    ws["L5"].font, ws["L5"].fill = font(WHITE, bold=True), FILL_HEAD
    tol, etol = "$B$3", "$B$4"
    R = lambda sheet, k, p: f"'{sheet}'!{COL[p]}{REG[(sheet, k)]}"  # noqa: E731
    eq = lambda a, b, t=tol: (lambda p: f"=ABS({a(p)}-{b(p)})<={t}")  # noqa: E731
    rows = []

    def chk(key, label, fn, periods=None, note=""):
        row = sh.line(key, label, fn, fmt="General", note=note, periods=periods)
        cells = [f"{COL[p]}{row}" for p in (periods or PERIODS)]
        ws[f"L{row}"] = f"=AND({','.join(cells)})"
        ws[f"L{row}"].font = font(bold=True)
        rows.append(row)

    sh.section("Balance sheet")
    chk("bs_balance", "Assets = liabilities + equity", lambda p: f"=ABS({R('BS', 'bs_diff', p)})<={tol}")
    for k, h, lab in (("ta", "total_assets", "Total assets"), ("tl", "total_liabilities", "Total liabilities"),
                      ("te", "total_equity", "Total equity")):
        chk(f"bs_{k}", f"{lab} = reported", eq(lambda p, k=k: R("BS", k, p), lambda p, h=h: R("Historicals", h, p)))
    sh.section("Income statement")
    for k, h, lab in (("gross_profit", "gross_profit", "Gross profit"), ("opex", "opex_total", "Operating expenses"),
                      ("ebit", "ebit", "EBIT"), ("nonop", "non_operating_total", "Non-operating total"),
                      ("pbt", "pbt", "Profit before tax"), ("net_income", "net_income", "Net income"),
                      ("ni_parent", "ni_parent", "Net income to parent")):
        chk(f"is_{k}", f"{lab} = reported", eq(lambda p, k=k: R("IS", k, p), lambda p, h=h: R("Historicals", h, p)))
    chk("is_eps", "EPS (computed) = reported basic EPS",
        eq(lambda p: R("IS", "eps", p), lambda p: R("Historicals", "eps_basic", p), etol),
        note="period-end shares vs weighted average")
    sh.section("Cash flow")
    chk("cf_pbt", "PBT in cash flow = PBT in income statement",
        eq(lambda p: R("CF", "pbt", p), lambda p: R("IS", "pbt", p)))
    chk("cf_change", "Net change in cash = reported",
        eq(lambda p: R("CF", "net_change", p), lambda p: R("Historicals", "net_change_cash", p)))
    chk("cf_cash", "Closing cash (CF) = cash on balance sheet",
        eq(lambda p: R("CF", "cash_close", p), lambda p: R("BS", "cash", p)))
    sh.section("Segments")
    chk("seg_ebit", "Segment profit total = EBIT",
        eq(lambda p: R("Segments", "op_total", p), lambda p: R("IS", "ebit", p)))
    chk("seg_others", "Unallocated revenue < 0.5% of group revenue",
        lambda p: f"=ABS({R('Segments', 'rev_others', p)})<0.005*{R('IS', 'revenue', p)}")
    sh.section("Source extraction (Python)")
    chk("py_checks", "All tie-outs in data/processed/checks.csv pass",
        lambda p: f"={R('Historicals', 'checks_passed', 'FY2021')}={R('Historicals', 'checks_total', 'FY2021')}",
        periods=["FY2021"], note="subtotals, restatements, monthly vs quarterly")
    sh.blank()
    r = sh.r
    ws[f"A{r}"] = "ALL CHECKS"
    ws[f"A{r}"].font = font(NAVY, bold=True, size=12)
    ws[f"L{r}"] = f"=AND({','.join(f'L{x}' for x in rows)})"
    ws[f"L{r}"].font = font(bold=True, size=12)
    REG[("Checks", "all")] = r
    tf_format(ws, f"C8:L{r}", "C8")
    return sh


def build_cover(wb, h):
    ws = wb.create_sheet("Cover", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 64
    ws["B2"] = "Delta Electronics, Inc. (2308.TW)"
    ws["B2"].font = font(NAVY, bold=True, size=18)
    ws["B3"] = "Initiation of coverage - financial model"
    ws["B3"].font = font(GREY, size=12)
    ws["B4"] = "Student research for educational purposes only. Not investment advice."
    ws["B4"].font = font(GREY, italic=True, size=9)
    items = [
        ("Company", "Delta Electronics, Inc.", None),
        ("Ticker / exchange", "2308 / Taiwan Stock Exchange", None),
        ("Units", "NT$ million unless stated (reported in NT$ thousands, Taiwan IFRS)", None),
        ("Latest reported period", "1H26 - six months to 30 June 2026 (reviewed)", None),
        ("Forecast base period", "9M26 once Q3 2026 results are published, ~29-30 Oct 2026 (decision D2)", None),
        ("Share price and date", "Set in M4 (yfinance, 2308.TW)", None),
        ("Rating", "[ANALYST TO REVIEW] - set in M5", None),
        ("12-month price target", "[ANALYST TO REVIEW] - set in M4/M5", None),
        ("Scenario", "='Assumptions'!B3", "link"),
        ("Model status", "M2 - historical model FY2021A-1H26A; forecasts added in M3", None),
        ("Last updated", dt.date.today().isoformat(), None),
    ]
    r = 6
    for k, v, style in items:
        ws[f"B{r}"], ws[f"C{r}"] = k, v
        ws[f"B{r}"].font = font(bold=True)
        ws[f"C{r}"].font = font(GREEN if style == "link" else BLACK)
        r += 1
    r += 1
    ws[f"B{r}"] = "ALL CHECKS PASS"
    ws[f"B{r}"].font = font(NAVY, bold=True, size=12)
    ws[f"C{r}"] = f"='Checks'!L{REG[('Checks', 'all')]}"
    ws[f"C{r}"].font = font(bold=True, size=12)
    ws[f"C{r}"].alignment = Alignment(horizontal="left")
    tf_format(ws, f"C{r}", f"C{r}")
    REG[("Cover", "checks")] = r
    r += 2
    ws[f"B{r}"] = "Key historicals (NT$ m)"
    ws[f"B{r}"].font = font(NAVY, bold=True)
    r += 1
    key_rows = [("Revenue", "IS", "revenue", F_NUM), ("EBIT margin", "IS", "ebit_margin", F_PCT),
                ("EPS, basic (NT$)", "IS", "eps", F_EPS), ("Free cash flow", "CF", "fcf", F_NUM),
                ("Net debt / (net cash)", "BS", "net_debt", F_NUM)]
    heads = ["FY2023", "FY2024", "FY2025", "1H26"]
    for i, p in enumerate(heads):
        c = ws.cell(row=r, column=4 + i, value=HEAD[p])
        c.font, c.fill, c.alignment = font(WHITE, bold=True), FILL_HEAD, Alignment(horizontal="right")
        ws.column_dimensions[c.column_letter].width = 12
    ws[f"B{r}"].fill = ws[f"C{r}"].fill = FILL_HEAD
    r += 1
    for label, sheet, key, fmt in key_rows:
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = font()
        for i, p in enumerate(heads):
            c = ws.cell(row=r, column=4 + i, value=f"='{sheet}'!{COL[p]}{REG[(sheet, key)]}")
            c.font, c.number_format = font(GREEN), fmt
        r += 1
    r += 1
    ws[f"B{r}"] = "Sheet guide"
    ws[f"B{r}"].font = font(NAVY, bold=True)
    r += 1
    guide = [("Assumptions", "Scenario selector; every forecast driver (M3) with its historical reference"),
             ("Historicals", "Extracted as-reported figures (inputs), FY2021A-FY2025A, 1H25A, 1H26A"),
             ("Segments", "Segment revenue and profit; AI data-centre anchors"),
             ("IS / BS / CF", "Three statements built by formula from Historicals"),
             ("Schedules", "Working capital, PP&E, debt and interest, tax, dividends, returns"),
             ("Checks", "Integrity tests; all must be TRUE")]
    for k, v in guide:
        ws[f"B{r}"], ws[f"C{r}"] = k, v
        ws[f"B{r}"].font, ws[f"C{r}"].font = font(bold=True), font()
        r += 1
    r += 1
    ws[f"B{r}"] = "Colour key"
    ws[f"B{r}"].font = font(NAVY, bold=True)
    r += 1
    for txt, col, fill in (("Blue: hard-coded input (sourced in data/sources.csv)", BLUE, None),
                           ("Black: formula", BLACK, None), ("Green: link to another sheet", GREEN, None),
                           ("Yellow fill: key input for the analyst", BLACK, FILL_KEY)):
        ws[f"B{r}"] = txt
        ws[f"B{r}"].font = font(col)
        if fill:
            ws[f"B{r}"].fill = fill
        r += 1
    ws.sheet_view.showGridLines = False


def main() -> int:
    h, segs, lines, checks = load()
    wb = Workbook()
    wb.remove(wb.active)
    build_historicals(wb, h, checks)
    build_is(wb)
    bs = build_bs(wb)
    build_cf(wb, bs)
    build_segments(wb, segs)
    build_schedules(wb)
    build_assumptions(wb)
    build_checks(wb)
    build_cover(wb, h)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
        ws.oddFooter.center.text = "Delta Electronics (2308.TW) model - &A - page &P of &N"
    order = ["Cover", "Assumptions", "Historicals", "Segments", "IS", "BS", "CF", "Schedules", "Checks"]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"Cover": NAVY, "Assumptions": "FFC000", "Checks": "70AD47"}.get(ws.title)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
